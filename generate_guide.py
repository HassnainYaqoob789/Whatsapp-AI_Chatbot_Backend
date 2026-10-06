import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Add Title
title = doc.add_heading('Naracord AI - Discord Moderation Setup Guide', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('This guide explains how to configure the Naracord AI Moderation Engine for your Discord server. Follow these step-by-step instructions to set up rules, permissions, and AI moderation logic.')

# Section 1
doc.add_heading('Step 1: General Configuration & Connection', level=1)
doc.add_paragraph('1. Connect your Discord Server: Ensure the bot is connected to your Discord server via the Naracord Dashboard.')
doc.add_paragraph('2. Integration Enabled: Turn this switch ON to activate the Discord connection.')
doc.add_paragraph('3. Allowed Channels (Listen In): Select the specific channels where you want the AI to moderate. Leave empty if you want the AI to monitor all text channels.')
doc.add_paragraph('4. Moderation Log Channel: Select a private channel (e.g., #mod-logs) where the bot will send reports when it punishes a user.')
doc.add_paragraph('5. AI Moderator Prompt (Custom Instructions): Give the AI a persona. E.g., "You are the official AI Moderator for our Server. Your job is to keep the community safe, friendly, and helpful."')

p = doc.add_paragraph()
r = p.add_run('[ INSERT IMAGE 1 HERE - Dashboard General Configuration ]')
r.bold = True
r.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')

# Section 2
doc.add_heading('Step 2: AI Moderation Engine & Permissions', level=1)
doc.add_paragraph('1. Moderation ON: Turn this ON to activate automated enforcement.')
doc.add_paragraph('2. Dry Run OFF: Keep this OFF for live enforcement. (If ON, the bot only logs violations but takes no action).')
doc.add_paragraph('3. AI Classification ON: Turn this ON to allow the AI to intelligently classify intent (e.g., detecting hidden bullying or disguised disrespect).')
doc.add_paragraph('4. Allowed Actions (Bot Permissions): Turn ON the actions you want the bot to be able to perform:')
doc.add_paragraph('   - Delete Messages: ON')
doc.add_paragraph('   - Warn Users: ON')
doc.add_paragraph('   - Timeout (Mute): ON')
doc.add_paragraph('   - Kick: ON')
doc.add_paragraph('   - Ban (Dangerous): OFF (Recommended)')

p = doc.add_paragraph()
r = p.add_run('[ INSERT IMAGE 2 HERE - Allowed Actions & Moderation Engine ]')
r.bold = True
r.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')

# Section 3
doc.add_heading('Step 3: Strike Escalation Policy & Spam Limits', level=1)
doc.add_paragraph('1. Spam Limit (Messages / 10s): Set to "5". If a user sends more than 5 messages in 10 seconds, they will get a warning strike for flooding.')
doc.add_paragraph('2. Strikes to Trigger: Set to "3". A user will be automatically punished after accumulating 3 warnings.')
doc.add_paragraph('3. Time Window (Days): Set to "7". A warning strike automatically expires from the user\'s record after 7 days.')
doc.add_paragraph('4. Action to Apply: Set to "Timeout". This is the punishment given when 3 strikes are reached.')
doc.add_paragraph('5. Timeout Duration (Mins): Set the duration (e.g., 60 for one hour).')

p = doc.add_paragraph()
r = p.add_run('[ INSERT IMAGE 3 HERE - Strike Escalation Policy ]')
r.bold = True
r.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')

# Section 4
doc.add_heading('Step 4: Dynamic Rules Configuration', level=1)
doc.add_paragraph('Add specific rules that the AI will enforce. Here are three recommended examples:')
doc.add_paragraph('Rule 1: Toxicity & Profanity')
doc.add_paragraph('- Description: "No swearing, hate speech, or offensive language."')
doc.add_paragraph('- Action: Delete')
doc.add_paragraph('- Keywords: Add bad words you explicitly want blocked.')
doc.add_paragraph('Rule 2: Self Promotion')
doc.add_paragraph('- Description: "No sharing of unauthorized promotional links or discord servers."')
doc.add_paragraph('- Action: Warn (Strike)')
doc.add_paragraph('- Keywords: "discord.gg", "http"')
doc.add_paragraph('Rule 3: Disrespect')
doc.add_paragraph('- Description: "Harassing, bullying, or disrespecting other community members."')
doc.add_paragraph('- Action: Timeout')

p = doc.add_paragraph()
r = p.add_run('[ INSERT IMAGE 4 HERE - Dynamic Rules ]')
r.bold = True
r.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')

# Section 5
doc.add_heading('Step 5: Community Terms & Privacy Policy', level=1)
doc.add_paragraph('Provide the bot with your server\'s Terms and Conditions. The bot will use these to answer user questions intelligently. For example:')
doc.add_paragraph('1. Respect everyone. Harassment, racism, and hate speech are strictly prohibited.')
doc.add_paragraph('2. Keep discussions on-topic in their respective channels.')
doc.add_paragraph('3. No NSFW (Not Safe For Work) content allowed.')
doc.add_paragraph('4. Listen to the AI Moderator and human staff.')

p = doc.add_paragraph()
r = p.add_run('[ INSERT IMAGE 5 HERE - Terms & Privacy Policy ]')
r.bold = True
r.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')

# Save Document
doc.save('C:/Users/DELL/Documents/GitHub/Whatsapp-Ai-Chatbot/Naracord_AI_Moderation_Guide.docx')

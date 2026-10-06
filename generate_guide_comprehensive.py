import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

def add_header(text, level=1):
    h = doc.add_heading(text, level=level)
    h.style.font.color.rgb = RGBColor(41, 128, 185)

def add_image_placeholder(filename, description):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'\n[ INSERT IMAGE HERE: {filename} ]\n({description})\n')
    run.bold = True
    run.font.color.rgb = RGBColor(231, 76, 60)

# Title
title = doc.add_heading('Naracord AI - Complete Setup & Moderation Guide', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Welcome to Naracord AI! This comprehensive document will guide you from A to Z on how to connect your Discord server, configure the AI brain, and set up a fully automated, intelligent moderation engine.')

# Section 1
add_header('Phase 1: Connecting Your Discord Server')
doc.add_paragraph('Before the AI can moderate your community, you must invite the bot to your Discord server.')
doc.add_paragraph('1. Go to your Naracord Dashboard and navigate to the "Discord Moderation" tab.')
doc.add_paragraph('2. Click on the blue "Connect to Discord" button.')
doc.add_paragraph('3. A popup will appear asking you to select your Discord Server. Choose your server from the dropdown list and click "Authorize".')
doc.add_paragraph('4. Important: Once the bot joins your server, open your Discord App, go to Server Settings -> Roles, and drag the "Naracord AI" role to the very top of the list. This ensures the bot has the authority to mute or kick rule-breakers.')

# Section 2
add_header('Phase 2: General Configuration')
doc.add_paragraph('Now that the bot is connected, let\'s configure where it operates.')
doc.add_paragraph('• Integration Enabled: Turn this switch ON to activate the bot on your server.')
doc.add_paragraph('• Allowed Channels (Listen In): By default, if you leave this empty, the AI will monitor ALL text channels. If you only want it to moderate specific channels (e.g., #general), select them here.')
doc.add_paragraph('• Moderation Log Channel: Select a private admin channel (e.g., #mod-logs). The AI will send detailed reports here whenever it deletes a message, issues a strike, or mutes a user.')

add_image_placeholder('media_1790768105941.png', 'Screenshot showing General Configuration section')

# Section 3
add_header('Phase 3: The AI Brain (Profile & Settings)')
doc.add_paragraph('Your AI is powered by OpenAI (GPT-4o). You need to give it a personality and instructions on how to behave.')
doc.add_paragraph('1. Navigate to the "Profile & Settings" tab on the left menu.')
doc.add_paragraph('2. AI System Prompt: This is the brain of your bot. Write clear instructions here. For example:')
doc.add_paragraph('   "You are the official AI Moderator for Minhaj\'s Server. Your job is to keep the community safe, friendly, and helpful. RULES: Be polite but strict against rule-breakers. Keep your answers concise (under 2 lines)."')
doc.add_paragraph('The AI will follow these instructions whenever a user asks it a question in the chat.')

add_image_placeholder('media_1790769650331.png', 'Screenshot showing Profile & AI Settings (System Prompt)')

# Section 4
add_header('Phase 4: AI Moderation Engine & Permissions')
doc.add_paragraph('This section controls how the AI enforces rules.')
doc.add_paragraph('• Moderation ON: Turn this ON to start actively enforcing rules.')
doc.add_paragraph('• Dry Run (Test Mode): If you turn this ON, the bot will only "pretend" to moderate. It will log violations in the mod-log channel but will NOT delete messages or punish users. Keep this OFF for live moderation.')
doc.add_paragraph('• AI Classification ON: This is a highly advanced feature. When turned ON, the AI doesn\'t just look for specific bad words; it actually reads the "intent" of the message. If someone is secretly bullying or harassing another user without using bad words, the AI will detect the toxic intent and punish them.')
doc.add_paragraph('• Allowed Actions: Turn ON the permissions you want to grant the bot (Delete Messages, Warn Users, Timeout, Kick). We recommend leaving "Ban" OFF to prevent accidental permanent bans.')

# Section 5
add_header('Phase 5: Strike Escalation & Spam Policy')
doc.add_paragraph('Naracord AI uses a "Rolling Strike" system, similar to YouTube\'s community guidelines.')
doc.add_paragraph('• Spam Limit: Set how many messages a user can send within 10 seconds (e.g., 5). If they exceed this, they get a spam strike.')
doc.add_paragraph('• Strikes to Trigger: How many warnings a user gets before a major punishment. (e.g., 3 strikes).')
doc.add_paragraph('• Time Window (Days): How long a strike stays on a user\'s record. If set to 7, a strike will automatically expire exactly 7 days after it was issued. Strikes expire individually, not all at once.')
doc.add_paragraph('• Action to Apply & Timeout Duration: If a user hits 3 strikes, what should happen? We recommend setting "Timeout" for "60" minutes. Note: After the 60 minutes are over, the user can chat again, but their strikes remain on their record until the 7 days are up. If they break a rule again, they will be timed out immediately.')

add_image_placeholder('media_1790768180043.png', 'Screenshot showing AI Moderation Engine, Spam Limit, and Strike Escalation')

# Section 6
add_header('Phase 6: Dynamic Rules')
doc.add_paragraph('Here you can create custom rules for your specific community.')
doc.add_paragraph('1. Rule Name: e.g., "Toxicity & Profanity" or "Self Promotion".')
doc.add_paragraph('2. Description (AI Context): Explain the rule to the AI so it knows what to look for (e.g., "No swearing or offensive language").')
doc.add_paragraph('3. Action: Choose what happens immediately when this rule is broken (Warn, Delete, Timeout, etc.).')
doc.add_paragraph('4. Trigger Keywords: (Optional) If you want specific words to ALWAYS trigger this rule instantly, add them here and press Enter (e.g., "discord.gg" to stop invite links).')

# Section 7
add_header('Phase 7: Community Terms & Privacy Policy')
doc.add_paragraph('The AI acts as an assistant. When community members ask questions like "What are the rules?" or "Can I post links?", the AI needs to know the answers.')
doc.add_paragraph('• Terms and Conditions: Paste your server\'s rules here (e.g., "1. Respect everyone. 2. No NSFW content."). The AI will memorize these and politely answer users who ask about them.')
doc.add_paragraph('• Privacy Policy: Paste your data guidelines here so the AI can reassure users about how their chat data is used (e.g., "We do not sell your data. The AI only scans messages for safety").')

add_image_placeholder('media_1790768351694.png', 'Screenshot showing Community Terms & Privacy Policy text areas')

doc.add_paragraph('\n\nCongratulations! Your Naracord AI system is now fully configured and ready to protect your community 24/7.')

# Save Document
doc.save('C:/Users/DELL/Documents/GitHub/Whatsapp-Ai-Chatbot/Naracord_Comprehensive_Client_Guide.docx')
